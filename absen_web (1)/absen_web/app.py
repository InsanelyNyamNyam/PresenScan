import functools
import hmac
import io
import os
import secrets
import sqlite3
import sys
import threading
from datetime import datetime
from time import sleep

from flask import Flask, abort, g, jsonify, render_template, request, send_file, session
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "students.db")        # sama seperti kode aslimu
EXCEL_PATH = os.path.join(BASE, "FileAbsen.xlsx")  # file Excel yang diisi saat scan

# =========================
# PASSWORD ADMIN  -> GANTI INI!
# (atau set lewat environment variable ADMIN_PASSWORD)
# =========================
DEFAULT_PASSWORD = "admin123"
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", DEFAULT_PASSWORD)

app = Flask(__name__)


def load_secret_key():
    """Kunci sesi login, dibuat sekali dan disimpan di secret.key."""
    path = os.path.join(BASE, "secret.key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(secrets.token_hex(32))
    with open(path) as f:
        return f.read().strip()


app.secret_key = load_secret_key()
app.config.update(SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_HTTPONLY=True)
excel_lock = threading.Lock()


def fill_row(sheet, absence_number, name, nis, class_name, date_display):
    """Logika sama seperti kode aslimu: baris = nomor absen + 1."""
    row = absence_number + 1
    values = [absence_number, name, nis, class_name, "HADIR", date_display]
    for col, value in enumerate(values, start=1):
        sheet.cell(row=row, column=col).value = value


# =========================
# DATABASE
# =========================

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_):
    db = g.pop("db", None)
    if db:
        db.close()


def init_db():
    """Tabel students sama persis dengan database_create.py.
    Tabel attendance baru: menyimpan riwayat absen (sumber data untuk Excel)."""
    con = sqlite3.connect(DB_PATH)
    con.executescript("""
    CREATE TABLE IF NOT EXISTS students (
        nis TEXT PRIMARY KEY,
        name TEXT NOT NULL,
        class_name TEXT NOT NULL,
        absence_number INTEGER NOT NULL
    );
    CREATE TABLE IF NOT EXISTS attendance (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nis TEXT NOT NULL REFERENCES students(nis),
        date TEXT NOT NULL,          -- format YYYY-MM-DD
        time TEXT NOT NULL,          -- format HH:MM:SS
        status TEXT NOT NULL DEFAULT 'HADIR',
        UNIQUE (nis, date)           -- 1 siswa hanya 1x absen per hari
    );
    """)
    con.commit()
    con.close()


# =========================
# ROUTES
# =========================

@app.get("/")
def index():
    return render_template("index.html")


@app.post("/api/scan")
def scan():
    nis = str((request.get_json(silent=True) or {}).get("nis", "")).strip()
    if not nis:
        return jsonify(status="invalid"), 400

    db = get_db()
    student = db.execute(
        "SELECT nis, name, class_name, absence_number FROM students WHERE nis = ?",
        (nis,),
    ).fetchone()

    if not student:
        return jsonify(status="not_found", nis=nis)

    status, time, excel_saved = mark_present(db, student)

    return jsonify(
        status=status,
        nis=student["nis"],
        name=student["name"],
        class_name=student["class_name"],
        absence_number=student["absence_number"],
        time=time,
        excel=excel_saved,
    )


def mark_present(db, student):
    """Catat HADIR hari ini ke database + Excel. Dipakai scan QR dan absen manual.
    Return (status, jam, excel_ok): status 'ok' atau 'duplicate'."""
    now = datetime.now()
    date, time = now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S")
    try:
        db.execute(
            "INSERT INTO attendance (nis, date, time) VALUES (?, ?, ?)",
            (student["nis"], date, time),
        )
        db.commit()
        return "ok", time, write_to_excel(student, now.strftime("%d/%m/%Y"))
    except sqlite3.IntegrityError:  # sudah absen hari ini
        row = db.execute(
            "SELECT time FROM attendance WHERE nis = ? AND date = ?",
            (student["nis"], date),
        ).fetchone()
        return "duplicate", row["time"], None


def write_to_excel(student, date_display):
    """Tulis 1 baris ke FileAbsen.xlsx langsung saat scan (seperti kode aslimu).
    Return False kalau gagal (misal file sedang dibuka di Excel)."""
    try:
        with excel_lock:
            workbook = load_workbook(EXCEL_PATH)
            fill_row(workbook.active, student["absence_number"], student["name"],
                     student["nis"], student["class_name"], date_display)
            workbook.save(EXCEL_PATH)
        return True
    except (PermissionError, OSError):
        return False


@app.get("/api/today")
def today():
    date = datetime.now().strftime("%Y-%m-%d")
    rows = get_db().execute(
        """SELECT s.absence_number, s.name, s.class_name, a.time
           FROM attendance a JOIN students s ON s.nis = a.nis
           WHERE a.date = ? ORDER BY a.time DESC LIMIT 100""",
        (date,),
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.get("/api/classes")
def classes():
    rows = get_db().execute(
        "SELECT DISTINCT class_name FROM students ORDER BY class_name"
    ).fetchall()
    return jsonify([r["class_name"] for r in rows])


@app.get("/download")
def download():
    class_name = request.args.get("class", "").strip()
    date = request.args.get("date") or datetime.now().strftime("%Y-%m-%d")
    if not class_name:
        abort(400, "Pilih kelas dulu")
    try:
        date_display = datetime.strptime(date, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        abort(400, "Format tanggal salah")

    rows = get_db().execute(
        """SELECT s.absence_number, s.name, s.nis, s.class_name
           FROM attendance a JOIN students s ON s.nis = a.nis
           WHERE a.date = ? AND s.class_name = ?""",
        (date, class_name),
    ).fetchall()

    with excel_lock:
        workbook = load_workbook(EXCEL_PATH)
    sheet = workbook.active
    for r in range(2, sheet.max_row + 1):      # kosongkan data lama (header tetap)
        for c in range(1, 7):
            sheet.cell(row=r, column=c).value = None
    for r in rows:
        fill_row(sheet, r["absence_number"], r["name"], r["nis"], r["class_name"], date_display)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = secure_filename(f"Absen_{class_name}_{date}.xlsx")
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# =========================
# ADMIN
# =========================

def admin_required(fn):
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("admin"):
            return jsonify(status="unauthorized"), 401
        return fn(*args, **kwargs)
    return wrapper


@app.get("/admin")
def admin_page():
    return render_template("admin.html", logged_in=bool(session.get("admin")))


@app.post("/admin/login")
def admin_login():
    password = str((request.get_json(silent=True) or {}).get("password", ""))
    if hmac.compare_digest(password.encode(), ADMIN_PASSWORD.encode()):
        session["admin"] = True
        return jsonify(status="ok")
    sleep(1)  # perlambat tebak-tebakan password
    return jsonify(status="wrong"), 401


@app.post("/admin/logout")
def admin_logout():
    session.pop("admin", None)
    return jsonify(status="ok")


@app.get("/api/admin/students")
@admin_required
def admin_students():
    date = datetime.now().strftime("%Y-%m-%d")
    rows = get_db().execute(
        """SELECT s.absence_number, s.name, s.nis, a.time
           FROM students s
           LEFT JOIN attendance a ON a.nis = s.nis AND a.date = ?
           WHERE s.class_name = ? ORDER BY s.absence_number""",
        (date, request.args.get("class", "")),
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.post("/api/admin/mark")
@admin_required
def admin_mark():
    nis = str((request.get_json(silent=True) or {}).get("nis", "")).strip()
    db = get_db()
    student = db.execute(
        "SELECT nis, name, class_name, absence_number FROM students WHERE nis = ?",
        (nis,),
    ).fetchone()
    if not student:
        return jsonify(status="not_found"), 404
    status, time, excel_ok = mark_present(db, student)
    return jsonify(status=status, name=student["name"], time=time, excel=excel_ok)


@app.post("/api/admin/reset")
@admin_required
def admin_reset():
    """Kosongkan FileAbsen.xlsx (header tetap) dan hapus absen HARI INI.
    Riwayat hari-hari sebelumnya di database tidak dihapus."""
    try:
        with excel_lock:
            workbook = load_workbook(EXCEL_PATH)
            sheet = workbook.active
            for r in range(2, sheet.max_row + 1):
                for c in range(1, 7):
                    sheet.cell(row=r, column=c).value = None
            workbook.save(EXCEL_PATH)
    except (PermissionError, OSError):
        return jsonify(status="excel_locked"), 409   # database belum disentuh
    db = get_db()
    cur = db.execute("DELETE FROM attendance WHERE date = ?",
                     (datetime.now().strftime("%Y-%m-%d"),))
    db.commit()
    return jsonify(status="ok", deleted=cur.rowcount)


init_db()

if __name__ == "__main__":
    # Kamera browser hanya jalan di localhost atau HTTPS.
    # Akses dari HP/perangkat lain lewat jaringan: jalankan `python app.py --https`
    if ADMIN_PASSWORD == DEFAULT_PASSWORD:
        print("\n[PERINGATAN] Password admin masih default ('admin123'). Ganti di app.py!\n")
    https = "--https" in sys.argv
    app.run(host="0.0.0.0", port=5000, ssl_context="adhoc" if https else None)
